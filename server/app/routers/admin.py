from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..db import get_db
from ..deps import get_current_staff
from ..models import (
    Checkin,
    ClassSession,
    Course,
    CourseSection,
    StudentSectionAssignment,
    User,
)
from ..schemas import (
    AdminCheckinOut,
    AdminClassSessionOut,
    AssignmentCreate,
    CheckinUpdate,
    ClassSessionCreate,
    CourseCreate,
    CourseOut,
    CourseUpdate,
    ImportResult,
    PointUpdateRequest,
    SectionCreate,
    SectionOut,
    SectionUpdate,
    StudentCreate,
    StudentOut,
    StudentSectionOut,
    StudentUpdate,
    TeacherOut,
    UserOut,
)

router = APIRouter()


def _get_manageable_session(
    db: Session, session_id: int, user: User
) -> tuple[ClassSession, CourseSection]:
    session = db.get(ClassSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="课程不存在")
    section = db.get(CourseSection, session.section_id)
    if section is None:
        raise HTTPException(status_code=404, detail="教学班不存在")
    if user.role == "teacher" and section.teacher_id != user.id:
        raise HTTPException(status_code=403, detail="只能管理自己的教学班")
    return session, section


def _to_out(
    db: Session, session: ClassSession, section: CourseSection, course: Course
) -> AdminClassSessionOut:
    checkin_count = db.scalar(
        select(func.count())
        .select_from(Checkin)
        .where(Checkin.session_id == session.id, Checkin.status == "valid")
    ) or 0
    total_students = db.scalar(
        select(func.count())
        .select_from(StudentSectionAssignment)
        .where(
            StudentSectionAssignment.section_id == session.section_id,
            StudentSectionAssignment.status == "active",
        )
    ) or 0
    return AdminClassSessionOut(
        id=session.id,
        course_name=course.course_name,
        section_name=section.section_name,
        section_code=section.section_code,
        class_date=session.class_date,
        start_time=session.start_time.strftime("%H:%M") if session.start_time else None,
        end_time=session.end_time.strftime("%H:%M") if session.end_time else None,
        classroom=session.classroom,
        checkin_status=session.checkin_status,
        class_status=session.class_status,
        latitude=float(session.latitude) if session.latitude is not None else None,
        longitude=float(session.longitude) if session.longitude is not None else None,
        distance_limit=float(session.distance_limit),
        checkin_count=checkin_count,
        total_students=total_students,
    )


def _load_out(db: Session, session: ClassSession) -> AdminClassSessionOut:
    section = db.get(CourseSection, session.section_id)
    course = db.get(Course, section.course_id)
    return _to_out(db, session, section, course)


@router.get("/me", response_model=UserOut)
def me(user: User = Depends(get_current_staff)):
    return UserOut.model_validate(user)


@router.get("/class-sessions", response_model=list[AdminClassSessionOut])
def list_class_sessions(
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    stmt = (
        select(ClassSession, CourseSection, Course)
        .join(CourseSection, ClassSession.section_id == CourseSection.id)
        .join(Course, CourseSection.course_id == Course.id)
        .order_by(ClassSession.class_date.desc(), ClassSession.start_time.desc())
        .limit(100)
    )
    if user.role == "teacher":
        stmt = stmt.where(CourseSection.teacher_id == user.id)

    rows = db.execute(stmt).all()
    return [_to_out(db, session, section, course) for session, section, course in rows]


@router.post("/class-sessions", response_model=AdminClassSessionOut)
def create_class_session(
    body: ClassSessionCreate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    section = db.get(CourseSection, body.section_id)
    if section is None:
        raise HTTPException(status_code=404, detail="教学班不存在")
    if user.role == "teacher" and section.teacher_id != user.id:
        raise HTTPException(status_code=403, detail="只能管理自己的教学班")
    if body.period_end < body.period_start:
        raise HTTPException(status_code=400, detail="结束节次不能小于开始节次")
    if (body.start_time is None) != (body.end_time is None):
        raise HTTPException(status_code=400, detail="开始时刻和结束时刻需同时填写")
    if body.start_time and body.end_time and body.end_time <= body.start_time:
        raise HTTPException(status_code=400, detail="结束时刻必须晚于开始时刻")
    if (body.latitude is None) != (body.longitude is None):
        raise HTTPException(status_code=400, detail="经纬度需同时填写")

    existing = db.scalar(
        select(ClassSession).where(
            ClassSession.section_id == section.id,
            ClassSession.class_date == body.class_date,
            ClassSession.period_start == body.period_start,
            ClassSession.period_end == body.period_end,
        )
    )
    if existing is not None:
        raise HTTPException(status_code=409, detail="该教学班同一时间已有课程")

    now = datetime.now()
    session = ClassSession(
        section_id=section.id,
        class_date=body.class_date,
        period_start=body.period_start,
        period_end=body.period_end,
        start_time=body.start_time,
        end_time=body.end_time,
        classroom=body.classroom,
        latitude=body.latitude,
        longitude=body.longitude,
        distance_limit=body.distance_limit,
        class_status="scheduled",
        checkin_status="draft",
        created_by=user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    return _load_out(db, session)


@router.post("/class-sessions/{session_id}/open", response_model=AdminClassSessionOut)
def open_checkin(
    session_id: int,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    session, section = _get_manageable_session(db, session_id, user)
    if session.class_status != "scheduled":
        raise HTTPException(status_code=409, detail="本节课已取消，不能开启签到")
    if session.latitude is None or session.longitude is None:
        raise HTTPException(status_code=400, detail="请先设置签到点")

    if session.checkin_status != "open":
        session.checkin_status = "open"
        session.opened_at = datetime.now()
        session.closed_at = None
        session.updated_at = datetime.now()
        db.commit()
        db.refresh(session)
    return _load_out(db, session)


@router.post("/class-sessions/{session_id}/close", response_model=AdminClassSessionOut)
def close_checkin(
    session_id: int,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    session, section = _get_manageable_session(db, session_id, user)
    if session.checkin_status != "closed":
        session.checkin_status = "closed"
        session.closed_at = datetime.now()
        session.updated_at = datetime.now()
        db.commit()
        db.refresh(session)
    return _load_out(db, session)


@router.put("/class-sessions/{session_id}/point", response_model=AdminClassSessionOut)
def update_point(
    session_id: int,
    body: PointUpdateRequest,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    session, section = _get_manageable_session(db, session_id, user)
    if session.checkin_status == "open":
        raise HTTPException(status_code=409, detail="签到已开启，请先关闭再修改点位")

    session.latitude = body.latitude
    session.longitude = body.longitude
    session.distance_limit = body.distance_limit
    session.updated_at = datetime.now()
    db.commit()
    db.refresh(session)
    return _load_out(db, session)


@router.get(
    "/class-sessions/{session_id}/checkins",
    response_model=list[AdminCheckinOut],
)
def list_checkins(
    session_id: int,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    _get_manageable_session(db, session_id, user)
    rows = db.scalars(
        select(Checkin)
        .where(Checkin.session_id == session_id)
        .order_by(Checkin.checked_in_at.asc())
    ).all()
    return [
        AdminCheckinOut(
            id=item.id,
            student_no=item.student_no_snapshot,
            student_name=item.student_name_snapshot,
            checked_in_at=item.checked_in_at,
            distance=float(item.distance),
            accuracy=float(item.accuracy) if item.accuracy is not None else None,
            status=item.status,
        )
        for item in rows
    ]


# ------------------------------------------------------------
# 课程 / 教学班 / 教师
# ------------------------------------------------------------

def _course_out(db: Session, course: Course) -> CourseOut:
    section_count = db.scalar(
        select(func.count())
        .select_from(CourseSection)
        .where(CourseSection.course_id == course.id)
    ) or 0
    return CourseOut(
        id=course.id,
        course_code=course.course_code,
        course_name=course.course_name,
        status=course.status,
        section_count=section_count,
    )


def _section_out(db: Session, section: CourseSection) -> SectionOut:
    course = db.get(Course, section.course_id)
    teacher = db.get(User, section.teacher_id) if section.teacher_id else None
    student_count = db.scalar(
        select(func.count())
        .select_from(StudentSectionAssignment)
        .where(
            StudentSectionAssignment.section_id == section.id,
            StudentSectionAssignment.status == "active",
        )
    ) or 0
    return SectionOut(
        id=section.id,
        course_id=section.course_id,
        course_code=course.course_code if course else "",
        course_name=course.course_name if course else "",
        section_code=section.section_code,
        section_name=section.section_name,
        teacher_id=section.teacher_id,
        teacher_name=teacher.name if teacher else None,
        status=section.status,
        student_count=student_count,
    )


def _ensure_teacher(db: Session, teacher_id: int) -> User:
    teacher = db.get(User, teacher_id)
    if teacher is None or teacher.role != "teacher" or teacher.status != "active":
        raise HTTPException(status_code=400, detail="任课教师不存在或不可用")
    return teacher


@router.get("/courses", response_model=list[CourseOut])
def list_courses(
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    courses = db.scalars(select(Course).order_by(Course.course_code.asc())).all()
    return [_course_out(db, course) for course in courses]


@router.post("/courses", response_model=CourseOut)
def create_course(
    body: CourseCreate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    code = body.course_code.strip()
    name = body.course_name.strip()
    if not code or not name:
        raise HTTPException(status_code=400, detail="课程编号和名称不能为空")
    if db.scalar(select(Course).where(Course.course_code == code)):
        raise HTTPException(status_code=409, detail="课程编号已存在")

    now = datetime.now()
    course = Course(
        course_code=code,
        course_name=name,
        status="active",
        created_at=now,
        updated_at=now,
    )
    db.add(course)
    db.commit()
    db.refresh(course)
    return _course_out(db, course)


@router.put("/courses/{course_id}", response_model=CourseOut)
def update_course(
    course_id: int,
    body: CourseUpdate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="课程不存在")

    if body.course_code is not None:
        code = body.course_code.strip()
        if not code:
            raise HTTPException(status_code=400, detail="课程编号不能为空")
        if code != course.course_code and db.scalar(
            select(Course).where(Course.course_code == code)
        ):
            raise HTTPException(status_code=409, detail="课程编号已存在")
        course.course_code = code
    if body.course_name is not None:
        name = body.course_name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="课程名称不能为空")
        course.course_name = name
    if body.status is not None:
        if body.status not in ("active", "disabled"):
            raise HTTPException(status_code=400, detail="课程状态不合法")
        course.status = body.status

    course.updated_at = datetime.now()
    db.commit()
    db.refresh(course)
    return _course_out(db, course)


@router.get("/sections", response_model=list[SectionOut])
def list_sections(
    course_id: int | None = None,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    stmt = select(CourseSection).order_by(
        CourseSection.course_id.asc(), CourseSection.section_code.asc()
    )
    if course_id is not None:
        stmt = stmt.where(CourseSection.course_id == course_id)
    return [_section_out(db, section) for section in db.scalars(stmt).all()]


@router.post("/sections", response_model=SectionOut)
def create_section(
    body: SectionCreate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    course = db.get(Course, body.course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="课程不存在")
    code = body.section_code.strip()
    name = body.section_name.strip()
    if not code or not name:
        raise HTTPException(status_code=400, detail="教学班编号和名称不能为空")
    if db.scalar(
        select(CourseSection).where(
            CourseSection.course_id == course.id,
            CourseSection.section_code == code,
        )
    ):
        raise HTTPException(status_code=409, detail="该课程下教学班编号已存在")
    if body.teacher_id is not None:
        _ensure_teacher(db, body.teacher_id)

    now = datetime.now()
    section = CourseSection(
        course_id=course.id,
        section_code=code,
        section_name=name,
        teacher_id=body.teacher_id,
        status="active",
        created_by=user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(section)
    db.commit()
    db.refresh(section)
    return _section_out(db, section)


@router.put("/sections/{section_id}", response_model=SectionOut)
def update_section(
    section_id: int,
    body: SectionUpdate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    section = db.get(CourseSection, section_id)
    if section is None:
        raise HTTPException(status_code=404, detail="教学班不存在")

    if body.section_code is not None:
        code = body.section_code.strip()
        if not code:
            raise HTTPException(status_code=400, detail="教学班编号不能为空")
        if code != section.section_code and db.scalar(
            select(CourseSection).where(
                CourseSection.course_id == section.course_id,
                CourseSection.section_code == code,
            )
        ):
            raise HTTPException(status_code=409, detail="该课程下教学班编号已存在")
        section.section_code = code
    if body.section_name is not None:
        name = body.section_name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="教学班名称不能为空")
        section.section_name = name
    if body.teacher_id is not None:
        _ensure_teacher(db, body.teacher_id)
        section.teacher_id = body.teacher_id
    if body.status is not None:
        if body.status not in ("active", "disabled"):
            raise HTTPException(status_code=400, detail="教学班状态不合法")
        section.status = body.status

    section.updated_at = datetime.now()
    db.commit()
    db.refresh(section)
    return _section_out(db, section)


@router.get("/teachers", response_model=list[TeacherOut])
def list_teachers(
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    teachers = db.scalars(
        select(User)
        .where(User.role == "teacher", User.status == "active")
        .order_by(User.name.asc())
    ).all()
    return [TeacherOut(id=item.id, name=item.name, phone=item.phone) for item in teachers]


# ------------------------------------------------------------
# 学生管理
# ------------------------------------------------------------

def _student_out(db: Session, student: User) -> StudentOut:
    rows = db.execute(
        select(StudentSectionAssignment, CourseSection, Course)
        .join(
            CourseSection,
            StudentSectionAssignment.section_id == CourseSection.id,
        )
        .join(Course, CourseSection.course_id == Course.id)
        .where(
            StudentSectionAssignment.student_id == student.id,
            StudentSectionAssignment.status == "active",
        )
        .order_by(Course.course_code.asc())
    ).all()
    sections = [
        StudentSectionOut(
            course_id=course.id,
            course_name=course.course_name,
            section_id=section.id,
            section_name=section.section_name,
        )
        for _, section, course in rows
    ]
    return StudentOut(
        id=student.id,
        student_no=student.student_no or "",
        name=student.name,
        class_name=student.class_name,
        status=student.status,
        sections=sections,
    )


@router.get("/students", response_model=list[StudentOut])
def list_students(
    keyword: str | None = None,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    stmt = select(User).where(User.role == "student")
    if keyword and keyword.strip():
        like = f"%{keyword.strip()}%"
        stmt = stmt.where(
            or_(
                User.student_no.like(like),
                User.name.like(like),
                User.class_name.like(like),
            )
        )
    students = db.scalars(stmt.order_by(User.student_no.asc())).all()
    return [_student_out(db, student) for student in students]


@router.post("/students", response_model=StudentOut)
def create_student(
    body: StudentCreate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    student_no = body.student_no.strip()
    name = body.name.strip()
    if len(student_no) != 10 or not student_no.isdigit():
        raise HTTPException(status_code=400, detail="学号必须是10位数字")
    if not name:
        raise HTTPException(status_code=400, detail="姓名不能为空")
    if db.scalar(select(User).where(User.student_no == student_no)):
        raise HTTPException(status_code=409, detail="学号已存在")

    now = datetime.now()
    student = User(
        role="student",
        student_no=student_no,
        name=name,
        class_name=body.class_name.strip() if body.class_name else None,
        status="active",
        created_at=now,
        updated_at=now,
    )
    db.add(student)
    db.commit()
    db.refresh(student)
    return _student_out(db, student)


@router.put("/students/{student_id}", response_model=StudentOut)
def update_student(
    student_id: int,
    body: StudentUpdate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    student = db.get(User, student_id)
    if student is None or student.role != "student":
        raise HTTPException(status_code=404, detail="学生不存在")

    if body.name is not None:
        name = body.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="姓名不能为空")
        student.name = name
    if body.class_name is not None:
        student.class_name = body.class_name.strip() or None
    if body.status is not None:
        if body.status not in ("active", "disabled"):
            raise HTTPException(status_code=400, detail="学生状态不合法")
        student.status = body.status

    student.updated_at = datetime.now()
    db.commit()
    db.refresh(student)
    return _student_out(db, student)


@router.post("/students/{student_id}/assignments", response_model=StudentOut)
def assign_student(
    student_id: int,
    body: AssignmentCreate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    student = db.get(User, student_id)
    if student is None or student.role != "student":
        raise HTTPException(status_code=404, detail="学生不存在")
    course = db.get(Course, body.course_id)
    section = db.get(CourseSection, body.section_id)
    if course is None or section is None or section.course_id != course.id:
        raise HTTPException(status_code=400, detail="课程和教学班不匹配")

    existing = db.scalar(
        select(StudentSectionAssignment).where(
            StudentSectionAssignment.student_id == student.id,
            StudentSectionAssignment.course_id == course.id,
        )
    )
    if existing is not None:
        if existing.section_id == section.id:
            return _student_out(db, student)
        raise HTTPException(status_code=409, detail="该学生在这门课已有教学班")

    now = datetime.now()
    db.add(
        StudentSectionAssignment(
            student_id=student.id,
            course_id=course.id,
            section_id=section.id,
            status="active",
            created_by=user.id,
            created_at=now,
            updated_at=now,
        )
    )
    db.commit()
    return _student_out(db, student)


def _cell_text(row, index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@router.post("/students/import", response_model=ImportResult)
async def import_students(
    file: UploadFile = File(...),
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 文件")
    content = await file.read()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel 解析失败") from exc

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Excel 没有数据")

    header = [str(value).strip() if value is not None else "" for value in rows[0]]

    def column(name: str) -> int:
        try:
            return header.index(name)
        except ValueError:
            return -1

    index_no = column("学号")
    index_name = column("姓名")
    index_class = column("班级")
    index_course = column("课程编号")
    index_section = column("教学班编号")
    if index_no < 0 or index_name < 0:
        raise HTTPException(status_code=400, detail="表头必须包含 学号、姓名")

    result = ImportResult()
    now = datetime.now()

    for row_index, row in enumerate(rows[1:], start=2):
        student_no = _cell_text(row, index_no)
        name = _cell_text(row, index_name)
        class_name = _cell_text(row, index_class)
        if not student_no and not name:
            continue
        if len(student_no) != 10 or not student_no.isdigit():
            result.errors.append(f"第{row_index}行：学号必须是10位数字")
            continue
        if not name:
            result.errors.append(f"第{row_index}行：姓名为空")
            continue

        student = db.scalar(select(User).where(User.student_no == student_no))
        changed = False
        if student is None:
            student = User(
                role="student",
                student_no=student_no,
                name=name,
                class_name=class_name or None,
                status="active",
                created_at=now,
                updated_at=now,
            )
            db.add(student)
            db.flush()
            result.created += 1
            changed = True
        else:
            if name and student.name != name:
                student.name = name
                changed = True
            if class_name and student.class_name != class_name:
                student.class_name = class_name
                changed = True
            if changed:
                student.updated_at = now
                result.updated += 1

        course_code = _cell_text(row, index_course)
        section_code = _cell_text(row, index_section)
        if course_code or section_code:
            if not course_code or not section_code:
                result.errors.append(f"第{row_index}行：课程编号和教学班编号需同时填写")
                continue
            course = db.scalar(select(Course).where(Course.course_code == course_code))
            section = (
                db.scalar(
                    select(CourseSection).where(
                        CourseSection.course_id == course.id,
                        CourseSection.section_code == section_code,
                    )
                )
                if course
                else None
            )
            if course is None or section is None:
                result.errors.append(f"第{row_index}行：课程或教学班不存在")
                continue

            existing = db.scalar(
                select(StudentSectionAssignment).where(
                    StudentSectionAssignment.student_id == student.id,
                    StudentSectionAssignment.course_id == course.id,
                )
            )
            if existing is None:
                db.add(
                    StudentSectionAssignment(
                        student_id=student.id,
                        course_id=course.id,
                        section_id=section.id,
                        status="active",
                        created_by=user.id,
                        created_at=now,
                        updated_at=now,
                    )
                )
                result.assigned += 1
            elif existing.section_id != section.id:
                result.errors.append(f"第{row_index}行：该学生在这门课已有其他教学班")
            elif not changed:
                result.skipped += 1

    db.commit()
    return result


# ------------------------------------------------------------
# 签到记录修改 / 导出
# ------------------------------------------------------------

def _checkin_out(item: Checkin) -> AdminCheckinOut:
    return AdminCheckinOut(
        id=item.id,
        student_no=item.student_no_snapshot,
        student_name=item.student_name_snapshot,
        checked_in_at=item.checked_in_at,
        distance=float(item.distance),
        accuracy=float(item.accuracy) if item.accuracy is not None else None,
        status=item.status,
    )


def _checkin_session(db: Session, checkin_id: int) -> tuple[Checkin, ClassSession]:
    item = db.get(Checkin, checkin_id)
    if item is None:
        raise HTTPException(status_code=404, detail="签到记录不存在")
    session = db.get(ClassSession, item.session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="课程不存在")
    return item, session


@router.put("/checkins/{checkin_id}", response_model=AdminCheckinOut)
def update_checkin(
    checkin_id: int,
    body: CheckinUpdate,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    item, session = _checkin_session(db, checkin_id)
    section = db.get(CourseSection, session.section_id)
    if user.role == "teacher" and section.teacher_id != user.id:
        raise HTTPException(status_code=403, detail="只能管理自己的教学班")

    if body.student_no_snapshot is not None:
        value = body.student_no_snapshot.strip()
        if value and (len(value) != 10 or not value.isdigit()):
            raise HTTPException(status_code=400, detail="学号必须是10位数字")
        item.student_no_snapshot = value
    if body.student_name_snapshot is not None:
        value = body.student_name_snapshot.strip()
        if not value:
            raise HTTPException(status_code=400, detail="姓名不能为空")
        item.student_name_snapshot = value
    if body.status is not None:
        if body.status not in ("valid", "deleted"):
            raise HTTPException(status_code=400, detail="签到状态不合法")
        item.status = body.status

    item.updated_at = datetime.now()
    db.commit()
    db.refresh(item)
    return _checkin_out(item)


@router.get("/class-sessions/{session_id}/checkins/export")
def export_checkins(
    session_id: int,
    user: User = Depends(get_current_staff),
    db: Session = Depends(get_db),
):
    session, section = _get_manageable_session(db, session_id, user)
    course = db.get(Course, section.course_id)
    rows = db.scalars(
        select(Checkin)
        .where(Checkin.session_id == session_id)
        .order_by(Checkin.checked_in_at.asc())
    ).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "签到记录"
    sheet["A1"] = "课程"
    sheet["B1"] = course.course_name if course else ""
    sheet["A2"] = "教学班"
    sheet["B2"] = section.section_name
    sheet["A3"] = "时间"
    sheet["B3"] = f"{session.class_date} {session.start_time or ''}-{session.end_time or ''}"
    sheet["A4"] = "教室"
    sheet["B4"] = session.classroom or ""

    headers = ["学号", "姓名", "签到时间", "距离(米)", "定位精度(米)", "状态"]
    for column_index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=6, column=column_index, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_index, item in enumerate(rows, start=7):
        sheet.cell(row=row_index, column=1, value=item.student_no_snapshot)
        sheet.cell(row=row_index, column=2, value=item.student_name_snapshot)
        time_cell = sheet.cell(row=row_index, column=3, value=item.checked_in_at)
        time_cell.number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(row=row_index, column=4, value=float(item.distance))
        sheet.cell(
            row=row_index,
            column=5,
            value=float(item.accuracy) if item.accuracy is not None else None,
        )
        sheet.cell(
            row=row_index,
            column=6,
            value="有效" if item.status == "valid" else "已删除",
        )

    widths = [14, 12, 22, 12, 14, 10]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column_index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"checkins_session_{session_id}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
